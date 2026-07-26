from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("research-agent-platform/docs/presentation-module-report.docx")


SECTIONS = [
    (
        "1. 模块定位",
        [
            "presentation 不是单纯把文字丢给模型再出 PPT，而是研究工作流中的最后一环之一。它要把实验方案、实验进展、结果图、论文草稿、审稿意见和阶段结论转成可直接汇报的材料。",
            "核心目标有两个：让用户快速形成可讲、可改、可复用的汇报内容；让生成结果保持证据可追溯，而不是只生成一套好看的幻灯片。",
        ],
    ),
    (
        "2. 需求与场景梳理",
        [
            "组会汇报是当前最优先场景。典型输入不是“我要做一个 PPT”，而是当前研究目标、本周实验、方案变动、结果图、问题、下周计划、汇报对象和时长。",
            "阶段汇报适合中期检查、项目推进和组内分享，重点是目标拆解、阶段结果、下一步计划和风险。",
            "论文汇报适合 paper talk、投稿前组内过稿和答辩预演，重点是问题定义、方法框架、核心实验结果、消融/对比和结论局限。",
            "Poster 与 PPT 逻辑相近但不是同一类产物。PPT 更适合讲述，Poster 更适合展示，需要更强的视觉压缩能力。",
            "未来还应支持中文/英文双语、不同场景模板、不同听众层级、时长适配和生成后的二次修改。",
        ],
    ),
    (
        "3. 功能需求",
        [
            "输入侧至少支持三类输入：手工结构化表单、文件上传、自然语言补充说明。建议字段包括标题、汇报类型、语言、听众、时长、实验目标、实验方案、实验进展、关键结果、图表文件、论文草稿或参考文献、问题、下周计划和重点强调项。",
            "输出侧至少要有 SLIDES_OUTLINE.md、SPEAKER_NOTES.md、QA_BRIEF.md、PPTX、逐页图片或预览图，以及可选 Poster 版本。",
            "质量要求是：内容必须和实验材料对齐，每页只表达一个主信息，关键数据能追溯到原始输入或图表，大纲生成后可人工确认，图片生成前不应跳过结构审查，输出要能继续编辑。",
        ],
    ),
    (
        "4. 推荐流程",
        [
            "组会标准流程：用户上传实验方案、进展、结果图、论文片段等材料；系统抽取结构化信息并生成汇报骨架；输出 SLIDES_OUTLINE.md 供用户确认；确认后生成逐页内容和页面图片；再组装 PPTX，同时输出讲稿和 Q&A brief；最后允许二次修改并导出终版。",
            "组会建议目录通常为：标题页、研究背景与目标、当前实验方案、已完成工作、核心结果、问题与风险、下一步计划、需要讨论的事项。",
            "Poster 不建议直接复用 PPT 页数，而应做压缩：保留问题、方法、结果、结论，删除过细过程性叙述，图表优先，文字尽量短。",
        ],
    ),
    (
        "5. 调试与联调建议",
        [
            "最小可用版本不要求一次性做满所有复杂能力，但至少要能接收结构化输入、汇总实验进展、生成汇报大纲、输出 PPT 草稿，并在人工确认后继续。",
            "调试重点包括：输入是否足够结构化、大纲是否贴合场景、每页是否信息过载、图和结论是否对应、讲稿是否和页面一致、最终文件能否正常打开。",
            "建议的测试层级是结构测试、内容测试、路由测试、生成测试和回归测试。",
        ],
    ),
    (
        "6. 对当前系统的实现建议",
        [
            "presentation 模块建议明确成分层流水线：input normalization、story extraction、outline generation、human approval、slide rendering、deck assembly、QA and export。",
            "这样做的好处是调试更清楚、中断后容易恢复、产物更适合追踪，后续扩展 Poster 也更容易。",
        ],
    ),
    (
        "7. Skill 化建议",
        [
            "这部分非常适合做成独立 skill，或者继续扩展现有 research-presentation-pipeline。它的流程稳定、产物稳定、输入模板稳定、复用性高，而且天然符合“先大纲、再图像、再成品”的固定工作流。",
            "建议补齐的能力包括：组会模板、Poster 模板、输入表单规范、大纲审查规则和 PPT 生成校验。",
        ],
    ),
    (
        "8. 结论",
        [
            "presentation 模块的关键不是会不会生成幻灯片，而是能不能把实验材料整理成可汇报、可审查、可复用的成果链。",
            "优先建议打通组会汇报标准路径：实验方案/进展上传 -> 结构化整理 -> 大纲确认 -> 页面生成 -> PPT 输出。Poster 可作为同一条流水线上的第二输出形态，在组会场景稳定后再展开。",
        ],
    ),
]


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Presentation 模块报告")
    set_font(r, size=20, bold=True, color="000000")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("面向科研 agent 的汇报与展示模块设计说明")
    set_font(r, size=11, color="555555")

    for heading, paras in SECTIONS:
        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(heading)
        set_font(r, size=16, bold=True, color="2E74B5")

        for text in paras:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_font(r, size=11)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("建议优先打通组会汇报标准路径：实验方案/进展上传 -> 结构化整理 -> 大纲确认 -> 页面生成 -> PPT 输出。")
    set_font(r, size=11, bold=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
