from __future__ import annotations

import hashlib
import ipaddress
import json
import re
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

from .models import UploadBatchRecord, WriteSourceConfig
from .publication_contracts import VenueRuleSource, VenueStyleCard, VenueStyleRule


TEXT_EXTENSIONS = {
    ".bib",
    ".cfg",
    ".csv",
    ".enw",
    ".html",
    ".ini",
    ".ipynb",
    ".json",
    ".log",
    ".md",
    ".nbib",
    ".out",
    ".py",
    ".r",
    ".rst",
    ".ris",
    ".sh",
    ".sql",
    ".tex",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
IMAGE_EXTENSIONS = {".gif", ".jpeg", ".jpg", ".png", ".svg", ".tif", ".tiff", ".webp"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".pptx", ".xlsx"} | IMAGE_EXTENSIONS
GENERATED_PAPER_FILES = {
    "PAPER_EVIDENCE_MAP.json",
    "PAPER_SELF_REVIEW.md",
    "PAPER_DELIVERY_REPORT.json",
    "CITATION_AUDIT.json",
}
IGNORED_CONTEXT_FILES = {
    "CLOUD_SYNC.json",
    "MANIFEST.md",
    "PAPER_SOURCE_SELECTION.json",
    "PAPER_EVIDENCE_METADATA.json",
    "PRESENTATION_SOURCE_SELECTION.json",
    "PRESENTATION_SOURCE_INDEX.md",
}
WORKSPACE_ORDER = (
    "paper",
    "figures",
    "results",
    "plan",
    "idea",
    "bib",
    "Content",
    "logs",
    "code",
    "wiki",
)

DISCIPLINE_CUES = {
    "ai_computer_science": ("artificial intelligence", "machine learning", "ai ", "benchmark", "llm", "模型", "算法", "计算机"),
    "medicine_clinical": ("clinical", "patient", "cohort", "randomized", "临床", "患者", "队列", "医学"),
    "environmental_science_engineering": ("environment", "hydrology", "watershed", "污染", "洪涝", "环境", "水文"),
    "chemistry_materials": ("material", "catalyst", "polymer", "化学", "材料", "催化", "表征"),
    "social_science": ("social science", "survey respondent", "causal inference", "社会科学", "问卷", "因果推断"),
    "law_legal": ("jurisdiction", "holding", "statute", "法律", "判例", "法学"),
    "arts_humanities": ("humanities", "archive", "translation", "人文", "文本细读", "译本"),
    "mathematics_theory": ("theorem", "lemma", "proof", "数学", "定理", "引理", "证明"),
    "life_science_biology": ("biology", "cell", "gene", "organism", "生物", "细胞", "基因"),
}


def select_discipline_profile(objective: str) -> dict:
    styles_path = Path(__file__).resolve().parents[2] / "resources" / "publication" / "styles" / "discipline-styles-v1.json"
    profiles = json.loads(styles_path.read_text(encoding="utf-8"))["profiles"]
    lowered = objective.lower()
    profile_id = next(
        (
            candidate
            for candidate, cues in DISCIPLINE_CUES.items()
            if any(cue in lowered for cue in cues)
        ),
        "general",
    )
    if profile_id == "general":
        return {
            "id": profile_id,
            "organization": "problem-evidence-analysis-boundary",
            "claim_calibration": "Match every substantive claim to supplied evidence and state material limits.",
            "paragraph_function": "Give each paragraph one analytical job and a visible transition.",
            "citation_placement": "Place citations beside the claim they support.",
            "reporting_conventions": ["source traceability", "conditions", "uncertainty", "limitations"],
        }
    return {"id": profile_id, **profiles[profile_id]}


def _document_profile_id(objective: str) -> str | None:
    if re.search(r"\b(?:phd|doctoral|master(?:'?s)?|dissertation|thesis)\b|博士|硕士|学位论文", objective, re.I):
        return "degree_thesis"
    return None


def _writing_genre(objective: str) -> str:
    if re.search(r"\b(?:systematic review|meta-analysis)\b|系统综述|元分析", objective, re.I):
        return "systematic_review"
    if re.search(r"\b(?:literature review|review article|systematic review|survey)\b|文献综述|综述文章", objective, re.I):
        return "literature_review"
    if re.search(r"\b(?:short communication|research letter|brief report)\b|短通讯|研究简报", objective, re.I):
        return "short_communication"
    if re.search(r"\b(?:perspective|commentary)\b|观点|评论", objective, re.I):
        return "perspective_commentary"
    if re.search(r"\b(?:method|protocol|resource)\b|方法学|协议|资源论文", objective, re.I):
        return "methods_protocol_resource"
    return "research_article"


def _writing_language_key(objective: str) -> str:
    if re.search(r"英文|英语|\benglish\b", objective, re.I):
        return "en"
    if re.search(r"中文|汉语|\bchinese\b", objective, re.I):
        return "zh"
    chinese_count = len(re.findall(r"[\u3400-\u9fff]", objective))
    latin_count = len(re.findall(r"[A-Za-z]", objective))
    return "zh" if chinese_count > latin_count else "en"


def _curated_exemplar_context(objective: str, discipline_id: str, language_key: str) -> list[str]:
    patterns_path = Path(__file__).resolve().parents[2] / "resources" / "publication" / "exemplars" / "curated-patterns-v1.json"
    try:
        patterns = json.loads(patterns_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []

    document_profile_id = _document_profile_id(objective)
    if document_profile_id:
        document_profile = patterns.get("document_profiles", {}).get(document_profile_id, {})
        rules = document_profile.get(language_key, [])
        if rules:
            return [
                f"Document profile: {document_profile_id}",
                "Curated exemplar patterns (derived guidance, not copied prose): " + " ".join(rules),
                "Never copy exemplar wording; source-specific submission rules still require current official guidance.",
            ]

    profile = patterns.get("profiles", {}).get(discipline_id, {})
    rules = profile.get(_writing_genre(objective), [])
    if not rules:
        return []
    return [
        "Curated exemplar patterns (derived guidance, not copied prose): " + " ".join(rules),
        "Never copy exemplar wording; source-specific submission rules still require current official guidance.",
    ]


def _style_corpus_context(objective: str, discipline_id: str, language_key: str) -> list[str]:
    corpus_root = Path(__file__).resolve().parents[2] / "resources" / "publication" / "style-corpus"
    matrices = []
    for filename in ("learning-matrix-v1.json",):
        try:
            matrices.append(json.loads((corpus_root / filename).read_text(encoding="utf-8")))
        except (OSError, ValueError):
            continue
    genre = _writing_genre(objective)
    matches = []
    for matrix in matrices:
        for item in matrix.get("rules", []):
            applies_to = set(item.get("applies_to", []))
            if not isinstance(item, dict) or not applies_to:
                continue
            if "all" in applies_to or genre in applies_to or discipline_id in applies_to or language_key in applies_to:
                matches.append(item)
    if not matches:
        return []
    rules = " ".join(str(item.get("rule", "")) for item in matches if item.get("rule"))
    boundaries = " ".join(str(item.get("boundary", "")) for item in matches if item.get("boundary"))
    source_ids = sorted({source for item in matches for source in item.get("source_ids", [])})
    return [
        "Verified style-corpus rules (derived from lawfully accessible full texts; not copied prose): " + rules,
        "Style-corpus boundaries: " + boundaries,
        "Style-corpus source IDs: " + ", ".join(source_ids),
    ]


def build_writing_style_context(objective: str, venue_card: VenueStyleCard | None = None) -> str:
    profile = select_discipline_profile(objective)
    language_key = _writing_language_key(objective)
    language = "Chinese" if language_key == "zh" else "English"
    rules = [
        f"Discipline profile: {profile['id']}",
        f"Organization: {profile['organization']}",
        f"Claim calibration: {profile['claim_calibration']}",
        f"Paragraph role: {profile['paragraph_function']}",
        f"Citation placement: {profile['citation_placement']}",
        "Reporting conventions: " + ", ".join(profile["reporting_conventions"]),
        f"Draft language: {language}. {profile.get('language_guidance', {}).get(language_key, '')}",
        "Use specific subjects and measured conditions instead of generic importance, novelty, smooth-transition, or capability language.",
        "Avoid these failure modes: " + "; ".join(profile.get("failure_modes", [])),
        "Preserve the user's terminology, evidence boundary, and section-specific purpose.",
    ]
    rules.extend(_curated_exemplar_context(objective, profile["id"], language_key))
    rules.extend(_style_corpus_context(objective, profile["id"], language_key))
    if venue_card:
        rules.append(f"Verified venue card: {venue_card.venue} ({venue_card.language})")
        rules.extend(
            f"Verified venue rule [{rule.category}]: {rule.instruction}"
            for rule in venue_card.rules
            if rule.status == "verified"
        )
        rules.extend(
            f"Author confirmation needed [{rule.category}]: {rule.instruction}"
            for rule in venue_card.rules
            if rule.status == "needs_author_confirmation"
        )
    return "\n".join(rules)


def load_workspace_venue_style_card(workspace_root: Path) -> VenueStyleCard | None:
    card_path = workspace_root / "Content" / "VENUE_STYLE_CARD.json"
    if not card_path.exists():
        return None
    try:
        return VenueStyleCard.model_validate_json(card_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def build_fulltext_venue_style_card(
    workspace_root: Path,
    *,
    venue: str,
    article_type: str,
    language: str,
    allowed_relative_paths: Iterable[str] | None = None,
    minimum_fulltexts: int = 3,
) -> VenueStyleCard | None:
    """Derive bounded, non-copying style rules from lawful same-session full texts."""
    allowed = set(allowed_relative_paths or ())
    paper_paths = sorted((workspace_root / "bib" / "papers").glob("*.pdf"))
    if allowed:
        paper_paths = [
            path for path in paper_paths
            if path.relative_to(workspace_root).as_posix() in allowed
        ]
    samples: list[tuple[Path, str]] = []
    for path in paper_paths:
        text = "\n".join(block for _, block, _ in _extract_source_blocks(path)).strip()
        if len(text) >= 2500:
            samples.append((path, text))
    if len(samples) < minimum_fulltexts:
        return None

    heading_counts: dict[str, int] = {}
    abstract_count = 0
    reference_count = 0
    for _, text in samples:
        lowered = text.lower()
        abstract_count += int("abstract" in lowered or "摘要" in text)
        reference_count += int("references" in lowered or "参考文献" in text)
        for heading in re.findall(r"(?m)^\s*(?:\d+(?:\.\d+)*\s+)?([A-Z][A-Za-z &/-]{2,70}|[\u4e00-\u9fff]{2,24})\s*$", text):
            normalized = re.sub(r"\s+", " ", heading).strip()
            heading_counts[normalized] = heading_counts.get(normalized, 0) + 1
    common_headings = [heading for heading, count in heading_counts.items() if count >= max(2, len(samples) // 2)]
    common_headings = common_headings[:12]
    source_urls = [f"workspace://{path.relative_to(workspace_root).as_posix()}" for path, _ in samples]
    rules = [
        VenueStyleRule(
            category="fulltext-structure",
            instruction=(
                "Use a reader-facing article structure consistent with the sampled full texts; adapt section names to the "
                "user's evidence and article type rather than copying any source wording."
                + (" Common observed headings: " + "; ".join(common_headings) + "." if common_headings else "")
            ),
            source_urls=source_urls,
        ),
        VenueStyleRule(
            category="fulltext-rhetoric",
            instruction=(
                f"Across {len(samples)} lawful full texts, keep abstracts and references as distinct article components "
                f"when appropriate (observed in {abstract_count}/{len(samples)} and {reference_count}/{len(samples)} samples). "
                "Learn organization and evidential moves only; never reuse source sentences."
            ),
            source_urls=source_urls,
        ),
    ]
    return VenueStyleCard(
        card_id=f"fulltext-{re.sub(r'[^a-z0-9]+', '-', venue.lower()).strip('-') or 'venue'}-{article_type}",
        venue=venue,
        language="zh" if language == "zh" else "en",
        sources=[VenueRuleSource(url=url, source_type="open_exemplar") for url in source_urls],
        rules=rules,
    )


def extract_venue_guidance_urls(objective: str) -> list[str]:
    urls: list[str] = []
    for candidate in re.findall(r"https://[^\s<>\]\[\)\}\"']+", objective):
        parsed = urlparse(candidate.rstrip(".,;；，。"))
        hostname = parsed.hostname or ""
        if parsed.scheme != "https" or not hostname or hostname.lower() == "localhost":
            continue
        try:
            if ipaddress.ip_address(hostname).is_private:
                continue
        except ValueError:
            pass
        if candidate not in urls:
            urls.append(candidate)
    return urls


def extract_writing_length_contract(objective: str) -> dict[str, int | str] | None:
    english_range = re.search(
        r"\b(\d{2,6})\s*(?:-|–|to)\s*(\d{2,6})\s*words?\b",
        objective,
        re.I,
    )
    if english_range:
        return {
            "minimum": int(english_range.group(1)),
            "maximum": int(english_range.group(2)),
            "unit": "words",
        }
    chinese_maximum = re.search(r"(?:不超过|至多|最多)\s*(\d{2,6})\s*字", objective)
    if chinese_maximum:
        return {"minimum": 0, "maximum": int(chinese_maximum.group(1)), "unit": "characters"}
    return None


def build_writing_length_guidance(objective: str) -> str:
    profiles_path = (
        Path(__file__).resolve().parents[2]
        / "resources"
        / "publication"
        / "length-profiles"
        / "article-section-lengths-v1.json"
    )
    try:
        payload = json.loads(profiles_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return ""

    if _document_profile_id(objective):
        profile_id = "degree_thesis"
    elif _writing_language_key(objective) == "zh" and _writing_genre(objective) == "literature_review":
        profile_id = "chinese_literature_review"
    elif _writing_language_key(objective) == "zh" and _writing_genre(objective) == "research_article":
        profile_id = "chinese_journal_article"
    else:
        profile_id = _writing_genre(objective)
    profile = payload.get("profiles", {}).get(profile_id)
    if not isinstance(profile, dict):
        return ""

    explicit = extract_writing_length_contract(objective)
    lines = [f"{profile['label']} fallback: {profile['default_total']}." ]
    if explicit:
        lines.append(
            "User-specified contract overrides profile defaults: "
            f"{explicit['minimum']}–{explicit['maximum']} {explicit['unit']}."
        )
        minimum = int(explicit["minimum"])
        maximum = int(explicit["maximum"])
        lines.append(
            f"Aim for approximately {(minimum + maximum) // 2} {explicit['unit']} before references, "
            "while preserving a verification margin inside the permitted range."
        )
    lines.append(
        "Section budget (references, tables, captions, and appendices are excluded unless the user or venue says otherwise):"
    )
    lines.extend(f"- {section['name']}: {section['budget']}" for section in profile.get("sections", []))
    if explicit and minimum > 0:
        target = (minimum + maximum) // 2
        scaled_budgets = _scaled_section_budgets(profile.get("sections", []), target, str(explicit["unit"]))
        if scaled_budgets:
            lines.append("Scaled section budget for this delivery (before references):")
            lines.extend(f"- {name}: {budget}" for name, budget in scaled_budgets)
    lines.append(str(profile.get("architecture", "")))
    lines.append(
        "Explicit user limits and verified target-venue instructions override these fallbacks. "
        "Treat percentages as planning bands, not quotas; do not add filler, and do not shorten by dropping evidence "
        "conditions, citations, uncertainty, or required section responsibilities."
    )
    return "\n".join(line for line in lines if line)


def _scaled_section_budgets(sections: list[dict], target: int, unit: str) -> list[tuple[str, str]]:
    scaled: list[tuple[str, str]] = []
    for section in sections:
        name = str(section.get("name", "")).strip()
        budget = str(section.get("budget", "")).strip()
        percentage = re.fullmatch(r"(\d+)–(\d+)%", budget)
        if not name or percentage is None:
            continue
        lower = round(target * int(percentage.group(1)) / 100)
        upper = round(target * int(percentage.group(2)) / 100)
        scaled.append((name, f"{lower}–{upper} {unit}"))
    return scaled


def assess_writing_length(text: str, contract: dict[str, int | str]) -> dict[str, int | str | bool]:
    unit = str(contract["unit"])
    if unit == "characters":
        count = len(re.sub(r"\s", "", text))
    else:
        count = len(re.findall(r"\b[\w'-]+\b", text))
    return {
        "unit": unit,
        "count": count,
        "valid": int(contract["minimum"]) <= count <= int(contract["maximum"]),
    }


_INLINE_MATERIAL_LABEL = re.compile(
    r"(?:以下(?:研究)?材料|(?:研究)?材料如下|研究计划如下|研究结果如下|原文如下|稿件如下|草稿如下|"
    r"(?:research\s+)?materials?\s*(?:are|is|below)|(?:research\s+)?plan\s*(?:is|below)|"
    r"(?:manuscript|draft)\s*(?:is|below))(?:[^\n:：]{0,200})(?:\s*(?:：|:)|\n)",
    re.IGNORECASE,
)


def extract_inline_write_material(objective: str, *, minimum_characters: int = 80) -> str:
    """Return explicitly labelled writing material embedded in a /write request.

    A topic-led request must remain topic-led: material is accepted only after an
    unambiguous user label and only when it has enough substantive content to be
    a factual source.  This keeps ordinary instructions such as "write an
    introduction about …" from accidentally bypassing retrieval.
    """
    match = _INLINE_MATERIAL_LABEL.search(objective)
    if match is None:
        return ""
    material = objective[match.end() :].strip()
    # A label followed by a very short phrase is normally still an instruction,
    # not a source package.  Count visible characters so Chinese material is
    # handled the same way as space-delimited English prose.
    visible = re.sub(r"\s+", "", material)
    if len(visible) < minimum_characters:
        return ""
    return material


def has_explicit_write_source(objective: str, workspace_root: Path) -> bool:
    """Whether the user explicitly chose source scope or workspace paths."""
    return _requested_source_scope(objective) != "auto" or bool(
        _extract_explicit_refs(objective, workspace_root)
    )


def resolve_write_source_config(
    objective: str,
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    *,
    source_limit: int = 50,
) -> WriteSourceConfig:
    requested_scope = _requested_source_scope(objective)
    explicit_refs = _extract_explicit_refs(objective, workspace_root)
    batches = list(upload_batches)
    latest_batch = batches[-1] if batches else None
    latest_refs = _valid_refs(
        workspace_root,
        latest_batch.relative_paths if latest_batch else (),
    )[:source_limit]
    workspace_refs = select_write_workspace_sources(workspace_root, limit=source_limit)

    if explicit_refs:
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="selected",
            source_refs=_expand_refs(workspace_root, explicit_refs, source_limit),
            selection_reason="用户在 /write 指令中明确指定了写作材料。",
        )
    if requested_scope == "attachments":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="attachments",
            source_refs=latest_refs,
            upload_batch_id=latest_batch.upload_batch_id if latest_batch else "",
            selection_reason="用户要求仅使用最近一次上传批次。",
        )
    if requested_scope == "session":
        session_refs = _deduplicate(
            ref
            for batch in batches
            for ref in _valid_refs(workspace_root, batch.relative_paths)
        )[:source_limit]
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="session",
            source_refs=session_refs,
            upload_batch_id=latest_batch.upload_batch_id if latest_batch else "",
            selection_reason="用户要求使用当前会话上传的全部写作材料。",
        )
    if requested_scope == "workspace":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="workspace",
            source_refs=workspace_refs,
            selection_reason="用户要求检索工作区内的论文、实验、图表和研究上下文。",
        )
    if requested_scope == "selected":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="selected",
            source_refs=[],
            selection_reason="用户要求使用指定材料，但指令中没有解析到有效路径。",
        )

    combined = _deduplicate([*latest_refs, *workspace_refs])[:source_limit]
    if latest_refs:
        reason = "自动模式优先纳入最近上传材料，并补充工作区内相关实验、图表、计划和文献证据。"
        scope = "session"
    else:
        reason = "未发现本次上传或显式路径，自动冻结工作区内与论文写作相关的材料。"
        scope = "workspace"
    return WriteSourceConfig(
        requested_scope="auto",
        resolved_scope=scope,
        source_refs=combined,
        upload_batch_id=latest_batch.upload_batch_id if latest_batch and latest_refs else "",
        selection_reason=reason,
    )


def select_write_workspace_sources(workspace_root: Path, *, limit: int = 50) -> list[str]:
    ranked: list[tuple[int, int, str]] = []
    for directory_rank, directory_name in enumerate(WORKSPACE_ORDER):
        directory = workspace_root / directory_name
        if not directory.exists():
            continue
        for path in directory.rglob("*"):
            if not _is_source_file(path):
                continue
            relative = path.relative_to(workspace_root).as_posix()
            priority = directory_rank * 100
            lowered = path.stem.lower()
            if directory_name == "paper" and path.parent.name == "uploads":
                priority -= 70
            if directory_name == "paper" and any(
                term in lowered for term in ("final", "revised", "draft", "终稿", "定稿", "修订稿", "草稿")
            ):
                priority -= 50
            if directory_name == "figures" and path.suffix.lower() in IMAGE_EXTENSIONS:
                priority -= 30
            ranked.append((priority, -path.stat().st_mtime_ns, relative))
    return [relative for _, _, relative in sorted(ranked)[:limit]]


def collect_paper_evidence(
    workspace_root: Path,
    source_refs: Iterable[str],
    *,
    total_limit: int = 40000,
) -> list[dict]:
    records: list[dict] = []
    used = 0
    for relative in source_refs:
        source = workspace_root / relative
        if not _is_source_file(source):
            continue
        suffix = source.suffix.lower()
        extracted = _extract_source_blocks(source)
        if not extracted and suffix in IMAGE_EXTENSIONS:
            extracted = [(None, "", "visual")]
        for page, text, evidence_type in extracted:
            excerpt = text.strip()[:6000]
            if excerpt and used + len(excerpt) > total_limit and records:
                break
            identity = f"{relative}:{page or 0}:{evidence_type}:{excerpt[:500]}"
            record = {
                "evidence_id": "PE-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:10].upper(),
                "source_path": relative,
                "page": page,
                "evidence_type": evidence_type,
                "provenance": "EXTRACTED" if excerpt else "SOURCE_FILE",
                "excerpt": excerpt,
            }
            records.append(record)
            used += len(excerpt)
            if len(records) >= 100 or used >= total_limit:
                return records
    return records


SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "摘要": ("摘要", "abstract"),
    "关键词": ("关键词", "keywords"),
    "引言": ("引言", "introduction"),
    "文献综述": ("文献综述", "related work", "literature review"),
    "方法": ("研究方法", "方法", "methods", "methodology", "materials and methods"),
    "结果": ("研究结果", "结果", "experiments", "results"),
    "讨论": ("讨论", "discussion"),
    "结论": ("结论", "conclusion"),
}


def infer_source_section(workspace_root: Path, source_refs: Iterable[str]) -> dict[str, str | float]:
    """Infer a pasted/uploaded manuscript section without overriding an explicit request."""
    scores = {section: 0 for section in SECTION_ALIASES}
    evidence: dict[str, str] = {}
    scanned = 0
    for relative in source_refs:
        source = workspace_root / relative
        if not _is_source_file(source):
            continue
        for _, text, _ in _extract_source_blocks(source):
            excerpt = text[:12000]
            scanned += len(excerpt)
            for section, aliases in SECTION_ALIASES.items():
                for alias in aliases:
                    if re.search(rf"(?im)^\s*(?:#+\s*)?{re.escape(alias)}\s*$", excerpt):
                        scores[section] += 6
                        evidence.setdefault(section, f"heading `{alias}`")
                    elif re.search(rf"(?i)\b{re.escape(alias)}\b", excerpt):
                        scores[section] += 1
                        evidence.setdefault(section, f"section cue `{alias}`")
            if scanned >= 24000:
                break
        if scanned >= 24000:
            break
    ranked = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    section, score = ranked[0]
    runner_up = ranked[1][1] if len(ranked) > 1 else 0
    confidence = "high" if score >= 6 and score >= runner_up + 3 else "medium" if score >= 3 else "low"
    return {
        "section": section if score else "",
        "confidence": confidence,
        "reason": evidence.get(section, "no reliable section heading or cue found"),
    }


def paper_evidence_to_json(records: Iterable[dict]) -> str:
    return json.dumps(list(records), ensure_ascii=False, indent=2)


def paper_evidence_prompt(records: Iterable[dict], *, limit: int = 30000) -> str:
    blocks: list[str] = []
    used = 0
    for record in records:
        location = record["source_path"]
        if record.get("page"):
            location += f" page {record['page']}"
        excerpt = record.get("excerpt") or "[visual or non-text source]"
        block = f"[{record['evidence_id']}] {location} | {record['evidence_type']} | {record['provenance']}\n{excerpt}"
        if used + len(block) > limit and blocks:
            break
        blocks.append(block)
        used += len(block)
    return "\n\n".join(blocks)


def select_venue_profile(objective: str) -> dict:
    lowered = objective.lower()
    if any(
        term in lowered
        for term in (
            "学位论文",
            "博士论文",
            "硕士论文",
            "博士学位",
            "硕士学位",
            "thesis",
            "dissertation",
            "章节润色",
            "章节修改",
            "论文框架",
        )
    ):
        name = "degree-thesis-section"
        required = ["结构诊断", "可替换研究框架", "未解决作者输入"]
    elif any(term in lowered for term in ("综述", "survey", "review article", "narrative review")):
        name = "review-article"
        required = ["Introduction", "Review Methodology", "Synthesis", "Limitations", "Conclusion", "References"]
    else:
        name = "research-article"
        required = ["Abstract", "Introduction", "Method", "Experiments or Results", "Limitations", "Conclusion", "References"]
    venue = next((item.upper() for item in ("ieee", "acm") if item in lowered), "")
    if "nature" in lowered:
        venue = "Nature"
    return {
        "profile": name,
        "venue": venue or "unspecified",
        "required_sections": required,
        "note": "Content requirements are validated independently from publisher template compliance.",
    }


def build_paper_quality_reports(
    workspace_root: Path,
    manuscript_relative: str,
    evidence_records: Iterable[dict],
    venue_profile: dict,
    *,
    compile_status_relative: str = "",
) -> tuple[dict, dict]:
    evidence_records = list(evidence_records)
    manuscript_path = workspace_root / manuscript_relative
    manuscript = manuscript_path.read_text(encoding="utf-8", errors="ignore") if manuscript_path.exists() else ""
    bib_keys = _workspace_bib_keys(workspace_root)
    manuscript_body = _markdown_body_before_references(manuscript)
    cited_keys = _cited_keys(manuscript_body)
    all_citation_keys = _cited_keys(manuscript)
    unknown_keys = sorted(set(all_citation_keys) - bib_keys)
    uncited_keys = sorted(bib_keys - set(cited_keys))
    body_citation_coverage = round(len(set(cited_keys) & bib_keys) / len(bib_keys), 3) if bib_keys else None
    corpus_coverage_status = (
        "not_applicable"
        if not bib_keys
        else "pass" if not uncited_keys else "needs_attention"
    )
    citation_audit = {
        "manuscript": manuscript_relative,
        "cited_keys": cited_keys,
        "known_bibliography_keys": sorted(bib_keys),
        "unknown_keys": unknown_keys,
        "uncited_keys": uncited_keys,
        "body_citation_coverage": body_citation_coverage,
        "corpus_coverage_status": corpus_coverage_status,
        "status": "pass" if not unknown_keys and corpus_coverage_status != "needs_attention" else "needs_attention",
        "note": (
            "Key resolution and body-use coverage are deterministic; semantic citation support still requires "
            "author verification. A reference-list entry alone does not count as support for a manuscript claim."
        ),
    }

    headings = _markdown_headings(manuscript)
    missing_sections = [
        requirement
        for requirement in venue_profile.get("required_sections", [])
        if not _section_present(requirement, headings)
    ]
    placeholders = re.findall(
        r"AUTHOR INPUT NEEDED|CITATION NEEDED|PLACEHOLDER|\bTODO\b|待补充|待补|需要作者",
        manuscript,
        re.I,
    )
    known_evidence_ids = {str(item.get("evidence_id", "")) for item in evidence_records}
    referenced_evidence_ids = sorted(set(re.findall(r"\bPE-[A-F0-9]{10}\b", manuscript, re.I)))
    unknown_evidence_ids = sorted(set(referenced_evidence_ids) - known_evidence_ids)
    compile_status = "not_requested"
    if compile_status_relative:
        compile_path = workspace_root / compile_status_relative
        if compile_path.exists():
            status_text = compile_path.read_text(encoding="utf-8", errors="ignore").lower()
            compile_status = "success" if "status: success" in status_text else "unavailable_or_failed"

    findings = [
        _finding("manuscript_exists", bool(manuscript.strip()), "hard", manuscript_relative),
        _finding("source_set_nonempty", bool(evidence_records), "soft", f"{len(evidence_records)} evidence records"),
        _finding("required_sections", not missing_sections, "hard", "Missing: " + ", ".join(missing_sections) if missing_sections else "Required sections present"),
        _finding("citation_keys_resolve", not unknown_keys, "hard", "Unknown: " + ", ".join(unknown_keys) if unknown_keys else f"{len(cited_keys)} explicit citation keys checked"),
        _finding("evidence_ids_resolve", not unknown_evidence_ids, "hard", "Unknown: " + ", ".join(unknown_evidence_ids) if unknown_evidence_ids else f"{len(referenced_evidence_ids)} evidence IDs checked"),
        _finding("no_placeholders", not placeholders, "soft", f"{len(placeholders)} unresolved placeholders"),
        {
            "check": "latex_compile",
            "status": compile_status,
            "severity": "soft",
            "detail": compile_status_relative or "No LaTeX compile status requested.",
        },
    ]
    hard_violations = [item for item in findings if item["severity"] == "hard" and item["status"] == "violated"]
    delivery = {
        "manuscript": manuscript_relative,
        "profile": venue_profile,
        "status": "pass" if not hard_violations else "needs_attention",
        "hard_violations": len(hard_violations),
        "findings": findings,
        "author_verification_required": True,
    }
    return citation_audit, delivery


def _requested_source_scope(objective: str) -> str:
    match = re.search(
        r"--(?:source|scope)(?:-scope)?(?:=|\s+)(auto|attachments|selected|session|workspace)\b",
        objective,
        re.I,
    )
    return match.group(1).lower() if match else "auto"


def _extract_explicit_refs(objective: str, workspace_root: Path) -> list[str]:
    candidates = re.findall(
        r"(?:`([^`]+)`|[\"']([^\"']+)[\"']|(?<![\w.-])((?:bib|plan|idea|code|figures|paper|rebuttal|wiki|Content|logs)[/\\][^\s,，;；]+))",
        objective,
        re.I,
    )
    values = [next((part for part in groups if part), "").rstrip(".,，。;；)") for groups in candidates]
    return _valid_refs(workspace_root, values)


def _valid_refs(workspace_root: Path, values: Iterable[str]) -> list[str]:
    root = workspace_root.resolve()
    refs: list[str] = []
    for value in values:
        normalized = str(value).replace("\\", "/").lstrip("./")
        target = (workspace_root / normalized).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            continue
        if target.exists() and normalized not in refs:
            refs.append(normalized)
    return refs


def _expand_refs(workspace_root: Path, refs: Iterable[str], limit: int) -> list[str]:
    expanded: list[str] = []
    for relative in refs:
        target = workspace_root / relative
        if _is_source_file(target):
            expanded.append(target.relative_to(workspace_root).as_posix())
        elif target.is_dir():
            expanded.extend(
                path.relative_to(workspace_root).as_posix()
                for path in sorted(target.rglob("*"))
                if _is_source_file(path)
            )
        if len(expanded) >= limit:
            break
    return _deduplicate(expanded)[:limit]


def _is_source_file(path: Path) -> bool:
    return (
        path.is_file()
        and path.suffix.lower() in DOCUMENT_EXTENSIONS
        and path.name not in GENERATED_PAPER_FILES | IGNORED_CONTEXT_FILES
        and ".compile" not in path.parts
    )


def _extract_source_blocks(path: Path) -> list[tuple[int | None, str, str]]:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXTENSIONS:
            return [(None, chunk, "text") for chunk in _chunks(path.read_text(encoding="utf-8", errors="ignore"), 5000)]
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            return [
                (page_number, (page.extract_text() or "")[:6000], "document_text")
                for page_number, page in enumerate(reader.pages[:30], start=1)
                if (page.extract_text() or "").strip()
            ]
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            text += "\n" + "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            return [(None, chunk, "document_text") for chunk in _chunks(text, 5000)]
        if suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(path)
            return [
                (
                    slide_number,
                    " | ".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()),
                    "slide_text",
                )
                for slide_number, slide in enumerate(deck.slides, start=1)
            ]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets[:8]:
                rows = [
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(min_row=1, max_row=100, values_only=True)
                ]
                blocks.append((None, f"Sheet: {sheet.title}\n" + "\n".join(rows), "table"))
            workbook.close()
            return blocks
    except Exception:
        return []
    return []


def _chunks(text: str, size: int) -> list[str]:
    cleaned = text.strip()
    return [cleaned[index : index + size] for index in range(0, len(cleaned), size)] if cleaned else []


def _workspace_bib_keys(workspace_root: Path) -> set[str]:
    keys: set[str] = set()
    for path in workspace_root.rglob("*.bib"):
        text = path.read_text(encoding="utf-8", errors="ignore")
        keys.update(re.findall(r"@\w+\s*\{\s*([^,\s]+)\s*,", text, re.I))
    return keys


def _cited_keys(text: str) -> list[str]:
    keys: list[str] = []
    for group in re.findall(r"\\cite\w*\{([^}]+)\}", text):
        keys.extend(key.strip() for key in re.split(r"[,;]", group) if key.strip())
    for citation_group in re.findall(r"\[([^\]]*@[\w:./-]+[^\]]*)\]", text):
        keys.extend(re.findall(r"@([\w:./-]+)", citation_group))
    return sorted(set(keys))


def _markdown_body_before_references(text: str) -> str:
    match = re.search(r"(?im)^#{1,6}\s+(?:references|bibliography|参考文献)\s*$", text)
    return text[: match.start()] if match else text


def _markdown_headings(text: str) -> list[str]:
    return [match.strip().lower() for match in re.findall(r"(?m)^#{1,6}\s+(.+?)\s*$", text)]


def _section_present(requirement: str, headings: list[str]) -> bool:
    aliases = {
        "experiments or results": ("experiment", "result", "evaluation", "实验", "结果"),
        "method": ("method", "methodology", "approach", "方法"),
        "review methodology": ("review methodology", "search strategy", "literature selection", "综述方法", "检索方法"),
        "synthesis": ("synthesis", "taxonomy", "comparative", "综合", "分类"),
        "limitations": ("limitation", "threat", "局限", "有效性威胁"),
        "references": ("reference", "bibliography", "参考文献"),
    }
    terms = aliases.get(requirement.lower(), (requirement.lower(),))
    return any(any(term in heading for term in terms) for heading in headings)


def _finding(check: str, passed: bool, severity: str, detail: str) -> dict:
    return {"check": check, "status": "pass" if passed else "violated", "severity": severity, "detail": detail}


def _deduplicate(values: Iterable[str]) -> list[str]:
    results: list[str] = []
    for value in values:
        if value not in results:
            results.append(value)
    return results
