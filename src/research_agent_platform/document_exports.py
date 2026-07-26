from __future__ import annotations

import io
import re
from dataclasses import dataclass
from xml.sax.saxutils import escape as xml_escape

try:
    from docx import Document
    from docx.shared import Pt
except Exception:  # pragma: no cover - optional dependency fallback
    Document = None
    Pt = None

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate, Spacer
except Exception:  # pragma: no cover - optional dependency fallback
    A4 = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    ListFlowable = None
    ListItem = None
    Paragraph = None
    Preformatted = None
    SimpleDocTemplate = None
    Spacer = None


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    items: list[str] | None = None


def detect_write_formats(objective: str) -> list[str]:
    lowered = objective.lower()
    formats: list[str] = []
    if any(token in lowered for token in ("latex", ".tex", " tex", "tex ")):
        formats.append("tex")
    if any(token in lowered for token in ("word", "docx", ".docx")):
        formats.append("docx")
    if "pdf" in lowered:
        formats.append("pdf")
    return formats or ["docx", "pdf"]


def parse_markdown(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    lines = markdown.replace("\r\n", "\n").splitlines()
    paragraph: list[str] = []
    bullets: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            text = " ".join(part.strip() for part in paragraph if part.strip()).strip()
            if text:
                blocks.append(Block(kind="paragraph", text=text))
            paragraph = []

    def flush_bullets() -> None:
        nonlocal bullets
        if bullets:
            blocks.append(Block(kind="bullets", items=bullets.copy()))
            bullets = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            flush_bullets()
            if in_code:
                blocks.append(Block(kind="code", text="\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped or stripped == "---":
            flush_paragraph()
            flush_bullets()
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            flush_bullets()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            blocks.append(Block(kind="heading", text=title, level=level))
            continue
        if stripped.startswith(("- ", "* ")):
            flush_paragraph()
            bullets.append(stripped[2:].strip())
            continue
        flush_bullets()
        paragraph.append(stripped)

    flush_paragraph()
    flush_bullets()
    if code_lines:
        blocks.append(Block(kind="code", text="\n".join(code_lines)))
    return blocks


def markdown_to_latex(markdown: str, title: str = "Research Draft") -> str:
    blocks = parse_markdown(markdown)
    lines = [
        r"\documentclass[11pt]{article}",
        r"\usepackage[margin=1in]{geometry}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{hyperref}",
        r"\usepackage{enumitem}",
        r"\title{" + _escape_latex(title) + "}",
        r"\date{}",
        r"\begin{document}",
        r"\maketitle",
        "",
    ]
    for block in blocks:
        if block.kind == "heading":
            command = {1: "section", 2: "subsection", 3: "subsubsection"}.get(block.level, "paragraph")
            lines.append(f"\\{command}{{{_escape_latex(block.text)}}}")
        elif block.kind == "paragraph":
            lines.append(_escape_latex(block.text))
            lines.append("")
        elif block.kind == "bullets":
            lines.append(r"\begin{itemize}[leftmargin=*]")
            for item in block.items or []:
                lines.append(r"\item " + _escape_latex(item))
            lines.append(r"\end{itemize}")
        elif block.kind == "code":
            lines.append(r"\begin{verbatim}")
            lines.append(block.text)
            lines.append(r"\end{verbatim}")
    lines.append(r"\end{document}")
    return "\n".join(lines) + "\n"


def markdown_to_docx_bytes(markdown: str) -> bytes:
    if Document is None or Pt is None:
        raise RuntimeError("python-docx is required for DOCX export")
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    blocks = parse_markdown(markdown)
    for block in blocks:
        if block.kind == "heading":
            level = min(max(block.level, 1), 4)
            document.add_heading(block.text, level=level)
        elif block.kind == "paragraph":
            document.add_paragraph(block.text)
        elif block.kind == "bullets":
            for item in block.items or []:
                document.add_paragraph(item, style="List Bullet")
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def markdown_to_pdf_bytes(markdown: str) -> bytes:
    if any(
        item is None
        for item in (
            A4,
            ParagraphStyle,
            getSampleStyleSheet,
            ListFlowable,
            ListItem,
            Paragraph,
            Preformatted,
            SimpleDocTemplate,
            Spacer,
        )
    ):
        raise RuntimeError("reportlab is required for PDF export")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    heading_styles = {
        1: styles["Heading1"],
        2: styles["Heading2"],
        3: styles["Heading3"],
    }
    body = ParagraphStyle(
        "BodyCN",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        spaceAfter=8,
    )
    code_style = ParagraphStyle(
        "CodeBlock",
        parent=styles["Code"],
        fontName="Courier",
        fontSize=8.5,
        leading=10.5,
        spaceAfter=8,
    )
    flow = []
    for block in parse_markdown(markdown):
        if block.kind == "heading":
            style = heading_styles.get(block.level, styles["Heading4"])
            flow.append(Paragraph(xml_escape(block.text), style))
        elif block.kind == "paragraph":
            flow.append(Paragraph(_paragraph_to_html(block.text), body))
        elif block.kind == "bullets":
            items = [ListItem(Paragraph(_paragraph_to_html(item), body)) for item in (block.items or [])]
            flow.append(ListFlowable(items, bulletType="bullet"))
        elif block.kind == "code":
            flow.append(Preformatted(block.text, code_style))
        flow.append(Spacer(1, 6))
    doc.build(flow)
    return buffer.getvalue()


def _paragraph_to_html(text: str) -> str:
    escaped = xml_escape(text)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"\*(.+?)\*", r"<i>\1</i>", escaped)
    return escaped


def _escape_latex(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    result = text
    for old, new in replacements.items():
        result = result.replace(old, new)
    return result
