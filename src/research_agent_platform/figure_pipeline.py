from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import UploadBatchRecord


DATA_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls"}
PREFERRED_DATA_ROOTS = (
    "figures/uploads",
    "Content/uploads",
    "paper/uploads",
    "figures",
    "Content",
    "paper",
)
TIME_TERMS = {
    "date",
    "datetime",
    "day",
    "epoch",
    "iteration",
    "month",
    "round",
    "step",
    "time",
    "timestamp",
    "week",
    "year",
    "日期",
    "时间",
    "年份",
    "轮次",
    "步数",
    "迭代",
}


class NoRenderableDataError(ValueError):
    def __init__(self, warnings: Iterable[str] = ()) -> None:
        self.warnings = list(warnings)
        super().__init__("No valid tabular data with numeric values was found.")


@dataclass
class TableData:
    source_ref: str
    sheet_name: str
    headers: list[str]
    rows: list[list[object]]


@dataclass
class CodeFigureResult:
    source_ref: str
    sheet_name: str
    chart_type: str
    x_column: str
    y_columns: list[str]
    row_count: int
    png_bytes: bytes
    svg_bytes: bytes
    warnings: list[str] = field(default_factory=list)


def render_code_figure(
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    *,
    objective: str = "",
) -> CodeFigureResult:
    warnings: list[str] = []
    for source_ref in select_data_sources(workspace_root, upload_batches):
        try:
            tables = read_tables(workspace_root / Path(source_ref), source_ref)
        except Exception as exc:
            warnings.append(f"{source_ref}: {exc.__class__.__name__}: {exc}")
            continue
        for table in tables:
            chart = _select_chart(table)
            if chart is None:
                warnings.append(f"{source_ref}: no sufficiently populated numeric column")
                continue
            chart_type, x_index, y_indices, plotted_rows = chart
            png_bytes, svg_bytes = _render_chart(
                table,
                chart_type=chart_type,
                x_index=x_index,
                y_indices=y_indices,
                plotted_rows=plotted_rows,
                objective=objective,
            )
            return CodeFigureResult(
                source_ref=source_ref,
                sheet_name=table.sheet_name,
                chart_type=chart_type,
                x_column=table.headers[x_index] if x_index is not None else "Row",
                y_columns=[table.headers[index] for index in y_indices],
                row_count=len(plotted_rows),
                png_bytes=png_bytes,
                svg_bytes=svg_bytes,
                warnings=warnings,
            )
    raise NoRenderableDataError(warnings)


def select_data_sources(
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
) -> list[str]:
    selected: list[str] = []
    seen: set[str] = set()

    def add(relative_path: str) -> None:
        normalized = relative_path.replace("\\", "/").lstrip("/")
        source = workspace_root / Path(normalized)
        if (
            normalized not in seen
            and source.is_file()
            and source.suffix.lower() in DATA_EXTENSIONS
            and "generated" not in {part.lower() for part in source.parts}
        ):
            selected.append(normalized)
            seen.add(normalized)

    batches = list(upload_batches)
    if batches:
        for relative_path in batches[-1].relative_paths:
            add(relative_path)
    for root_name in PREFERRED_DATA_ROOTS:
        root = workspace_root / Path(root_name)
        if root.exists():
            for source in sorted(root.rglob("*")):
                if source.is_file():
                    add(source.relative_to(workspace_root).as_posix())
    for source in sorted(workspace_root.rglob("*")):
        if source.is_file():
            add(source.relative_to(workspace_root).as_posix())
    return selected


def read_tables(path: Path, source_ref: str) -> list[TableData]:
    extension = path.suffix.lower()
    if extension in {".csv", ".tsv"}:
        return [_read_delimited_table(path, source_ref)]
    if extension == ".xlsx":
        return _read_xlsx_tables(path, source_ref)
    if extension == ".xls":
        raise ValueError("legacy .xls is not supported; save the workbook as .xlsx")
    raise ValueError(f"unsupported data extension: {extension}")


def _read_delimited_table(path: Path, source_ref: str) -> TableData:
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise ValueError("empty table")
    default_delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = default_delimiter
    records = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _table_from_records(source_ref, "", records)


def _read_xlsx_tables(path: Path, source_ref: str) -> list[TableData]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[TableData] = []
    try:
        for worksheet in workbook.worksheets:
            records = [list(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                tables.append(_table_from_records(source_ref, worksheet.title, records))
            except ValueError:
                continue
    finally:
        workbook.close()
    if not tables:
        raise ValueError("workbook has no non-empty table")
    return tables


def _table_from_records(
    source_ref: str,
    sheet_name: str,
    records: list[list[object]],
) -> TableData:
    records = [row for row in records if any(_clean_cell(value) != "" for value in row)]
    if len(records) < 3:
        raise ValueError("table needs a header and at least two data rows")
    width = max(len(row) for row in records)
    raw_headers = [_clean_cell(value) for value in records[0]] + [""] * width
    headers = [raw_headers[index] or f"Column {index + 1}" for index in range(width)]
    rows = [(row + [None] * width)[:width] for row in records[1:5001]]
    return TableData(source_ref=source_ref, sheet_name=sheet_name, headers=headers, rows=rows)


def _select_chart(
    table: TableData,
) -> tuple[str, int | None, list[int], list[list[object]]] | None:
    numeric_columns: list[int] = []
    for column_index in range(len(table.headers)):
        populated = [row[column_index] for row in table.rows if _clean_cell(row[column_index]) != ""]
        numeric_count = sum(_to_number(value) is not None for value in populated)
        if numeric_count >= 2 and numeric_count / max(1, len(populated)) >= 0.7:
            numeric_columns.append(column_index)
    if not numeric_columns:
        return None

    non_numeric_columns = [
        index for index in range(len(table.headers)) if index not in numeric_columns
    ]
    time_column = next(
        (
            index
            for index in range(len(table.headers))
            if _is_time_header(table.headers[index]) and index != numeric_columns[-1]
        ),
        None,
    )
    if time_column is not None:
        y_indices = [index for index in numeric_columns if index != time_column][:4]
        if y_indices:
            rows = _complete_rows(table.rows, time_column, y_indices)
            if len(rows) >= 2:
                return "line", time_column, y_indices, rows

    if non_numeric_columns:
        x_index = non_numeric_columns[0]
        y_indices = numeric_columns[:4]
        rows = _complete_rows(table.rows, x_index, y_indices)
        if len(rows) >= 2:
            return "bar", x_index, y_indices, rows[:30]

    if len(numeric_columns) >= 2:
        x_index, y_index = numeric_columns[:2]
        rows = _complete_rows(table.rows, x_index, [y_index])
        if len(rows) >= 2:
            return "scatter", x_index, [y_index], rows

    y_index = numeric_columns[0]
    rows = [row for row in table.rows if _to_number(row[y_index]) is not None]
    if len(rows) >= 2:
        return "line", None, [y_index], rows
    return None


def _render_chart(
    table: TableData,
    *,
    chart_type: str,
    x_index: int | None,
    y_indices: list[int],
    plotted_rows: list[list[object]],
    objective: str,
) -> tuple[bytes, bytes]:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "axes.edgecolor": "#B8BEC6",
            "axes.labelcolor": "#2F343B",
            "text.color": "#20242A",
        }
    )
    figure, axis = plt.subplots(figsize=(12, 7.5), dpi=160)
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")
    colors = ["#2463A7", "#D4553D", "#2D8A63", "#8A6BBE"]

    if x_index is None:
        x_values = list(range(1, len(plotted_rows) + 1))
        x_label = "Row"
    elif chart_type == "scatter":
        x_values = [_to_number(row[x_index]) for row in plotted_rows]
        x_label = table.headers[x_index]
    else:
        x_values = [_clean_cell(row[x_index]) for row in plotted_rows]
        x_label = table.headers[x_index]

    if chart_type == "bar":
        positions = list(range(len(plotted_rows)))
        group_width = 0.78 / len(y_indices)
        for series_index, column_index in enumerate(y_indices):
            offset = (series_index - (len(y_indices) - 1) / 2) * group_width
            values = [_to_number(row[column_index]) for row in plotted_rows]
            axis.bar(
                [position + offset for position in positions],
                values,
                width=group_width * 0.9,
                label=table.headers[column_index],
                color=colors[series_index % len(colors)],
            )
        axis.set_xticks(positions, x_values, rotation=35 if len(positions) > 8 else 0, ha="right" if len(positions) > 8 else "center")
    elif chart_type == "scatter":
        y_index = y_indices[0]
        y_values = [_to_number(row[y_index]) for row in plotted_rows]
        axis.scatter(x_values, y_values, s=48, color=colors[0], alpha=0.82, edgecolors="white", linewidths=0.6)
    else:
        for series_index, column_index in enumerate(y_indices):
            y_values = [_to_number(row[column_index]) for row in plotted_rows]
            axis.plot(
                x_values,
                y_values,
                marker="o",
                markersize=4.5,
                linewidth=2.2,
                label=table.headers[column_index],
                color=colors[series_index % len(colors)],
            )

    title = _chart_title(objective, table.source_ref)
    axis.set_title(title, loc="left", fontsize=18, fontweight="bold", pad=18)
    axis.set_xlabel(x_label, fontsize=11, labelpad=10)
    if len(y_indices) == 1:
        axis.set_ylabel(table.headers[y_indices[0]], fontsize=11, labelpad=10)
    axis.grid(axis="y", color="#E4E7EB", linewidth=0.8)
    axis.spines[["top", "right"]].set_visible(False)
    if len(y_indices) > 1 or chart_type == "line":
        axis.legend(frameon=False, loc="best")
    figure.text(0.01, 0.01, f"Source: {table.source_ref}", fontsize=8, color="#6B737C")
    figure.tight_layout(rect=(0, 0.035, 1, 1))

    png_buffer = io.BytesIO()
    svg_buffer = io.BytesIO()
    figure.savefig(png_buffer, format="png", bbox_inches="tight", facecolor="white")
    figure.savefig(svg_buffer, format="svg", bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return png_buffer.getvalue(), svg_buffer.getvalue()


def _complete_rows(
    rows: list[list[object]],
    x_index: int,
    y_indices: list[int],
) -> list[list[object]]:
    return [
        row
        for row in rows
        if _clean_cell(row[x_index]) != ""
        and all(_to_number(row[column_index]) is not None for column_index in y_indices)
    ]


def _to_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    percentage = text.endswith("%")
    if percentage:
        text = text[:-1].strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return number / 100 if percentage else number


def _clean_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def _is_time_header(header: str) -> bool:
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "", header.lower())
    return any(term in normalized for term in TIME_TERMS)


def _chart_title(objective: str, source_ref: str) -> str:
    cleaned = re.sub(r"^/(?:fig|figure)\s*", "", objective.strip(), flags=re.I)
    if cleaned and len(cleaned) <= 72:
        return cleaned
    return Path(source_ref).stem.replace("_", " ").strip().title() or "Research Result"
