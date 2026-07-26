from __future__ import annotations

import asyncio
import json
from pathlib import Path

from pptx import Presentation

from research_agent_platform import agent as agent_module
from research_agent_platform.agent import ResearchAgentService

from .conftest import run


def test_direct_chat_identity_reply(service: ResearchAgentService):
    result = run(service.chat(None, "你是谁"))
    assert result["status"] == "idle"
    assert result["task_id"] == ""
    assert result["checkpoint"] is None
    assert "科研" in result["text"]


def test_present_runs_without_routine_checkpoint(service: ResearchAgentService, isolated_env):
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    assert start["status"] == "completed"
    assert start["command"] == "/present"
    assert start["checkpoint"] is None
    assert Path(isolated_env["state_root"], "langgraph-checkpoints.pkl").exists()

    task = service.get_task(task_id)
    assert task is not None
    assert task.status == "completed"
    assert task.user_id == "local"
    artifact_root = Path(task.artifact_root).as_posix()
    assert f"local/{task.session_id}" in artifact_root
    assert Path(task.artifact_root, "presentation", "SLIDES_OUTLINE.md").exists()
    assert Path(task.artifact_root, "bib").is_dir()
    assert Path(task.artifact_root, "code").is_dir()
    assert Path(task.artifact_root, "figures").is_dir()
    assert Path(task.artifact_root, "paper").is_dir()
    assert Path(task.artifact_root, "Content").is_dir()
    assert len(task.approvals) == 0
    assert task.current_stage_name == "qa_brief"
    assert any(artifact.relative_path == "presentation/SLIDES_OUTLINE.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/SLIDE_CONTENT.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/SPEAKER_NOTES.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/QA_BRIEF.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/STAGE_REPORT.pptx" for artifact in task.artifacts)
    assert Path(task.artifact_root, "presentation", "slides", "SLIDE_01.png").exists()
    assert Path(task.artifact_root, "Content", "PRESENTATION_TEMPLATE.json").exists()
    notes_json = Path(task.artifact_root, "Content", "SPEAKER_NOTES.json")
    assert notes_json.exists()
    assert json.loads(notes_json.read_text(encoding="utf-8"))[0]["notes"]
    deck = Presentation(Path(task.artifact_root, "presentation", "STAGE_REPORT.pptx"))
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)


def test_paper_talk_uses_paper_delivery_name(service: ResearchAgentService):
    session = service.store.create_session()
    Path(session.workspace_root, "paper").mkdir(parents=True, exist_ok=True)
    Path(session.workspace_root, "paper", "FINAL_PAPER.md").write_text(
        "# Final Paper\n\nValidated result.", encoding="utf-8"
    )
    start = run(service.chat(session.session_id, "/present 论文汇报"))
    task = service.get_task(start["task_id"])
    assert task is not None
    assert Path(task.artifact_root, "presentation", "slides").exists()

    result = start
    task = service.get_task(task.task_id)

    assert result["status"] == "completed"
    assert task is not None
    assert any(artifact.relative_path == "presentation/PAPER_TALK.pptx" for artifact in task.artifacts)
    assert Path(task.artifact_root, "Content", "PRESENTATION_SOURCE_INDEX.md").read_text(
        encoding="utf-8"
    ).find("paper/FINAL_PAPER.md") >= 0


def test_present_source_set_is_frozen_at_task_start(service: ResearchAgentService):
    session = service.store.create_session()
    Path(session.workspace_root, "paper").mkdir(parents=True, exist_ok=True)
    Path(session.workspace_root, "paper", "INITIAL.md").write_text("initial", encoding="utf-8")

    start = run(service.chat(session.session_id, "/present --source workspace 论文汇报"))
    task = service.get_task(start["task_id"])
    assert task is not None
    assert task.presentation_source is not None
    assert "paper/INITIAL.md" in task.presentation_source.source_refs

    Path(session.workspace_root, "paper", "LATE.md").write_text("late", encoding="utf-8")
    assert start["status"] == "completed"
    selection = Path(
        task.artifact_root,
        "Content",
        "PRESENTATION_SOURCE_SELECTION.json",
    ).read_text(encoding="utf-8")

    assert "paper/INITIAL.md" in selection
    assert "paper/LATE.md" not in selection


def test_present_only_checkpoints_for_explicit_multiple_choices(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_with_choices(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Slides Outline" in user_prompt:
            return content + "\n\n## Decision Required\n- Option A: concise talk\n- Option B: detailed talk\n"
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_with_choices)
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    assert start["status"] == "waiting_human"
    assert start["checkpoint"]["title"] == "Presentation Outline Approval"
    rejected = run(service.reject_task(task_id, "请选择更详细的叙事路线"))
    task = service.get_task(task_id)

    assert rejected["status"] == "waiting_human"
    assert rejected["checkpoint"]["title"] == "Presentation Outline Approval"
    assert task is not None
    assert task.status == "waiting_human"
    assert len(task.approvals) == 2
    assert task.approvals[0].status == "rejected"
    assert task.approvals[1].status == "pending"
    assert any("打回" in item for item in task.progress_log)


def test_present_does_not_checkpoint_for_markdown_none_decision(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_with_none(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Slides Outline" in user_prompt:
            return content + "\n\n## Open Design Questions\n- Option A\n- Option B\n\n## Decision Required\n\n**None.**\n"
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_with_none)
    start = run(service.chat(None, "/present 做一份论文汇报"))

    assert start["status"] == "completed"
    assert start["checkpoint"] is None


def test_resume_after_restart(service: ResearchAgentService, isolated_env):
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    fresh_service = ResearchAgentService()
    task = fresh_service.get_task(task_id)

    assert start["status"] == "completed"
    assert task is not None
    assert task.status == "completed"
    assert Path(isolated_env["state_root"], "langgraph-checkpoints.pkl").exists()
    assert Path(task.artifact_root).exists()


def test_resume_after_interruption_during_presentation_delivery(service, monkeypatch):
    async def scenario():
        original_generate_image = agent_module.generate_image
        image_started = asyncio.Event()
        release_image = asyncio.Event()

        async def delayed_generate_image(**kwargs):
            image_started.set()
            await release_image.wait()
            return await original_generate_image(**kwargs)

        monkeypatch.setattr(agent_module, "generate_image", delayed_generate_image)
        running = asyncio.create_task(service.chat(None, "/present 做一个中文汇报"))
        await image_started.wait()
        tasks = service.list_tasks()
        assert len(tasks) == 1
        task_id = tasks[0].task_id
        running.cancel()
        try:
            await running
        except asyncio.CancelledError:
            service.record_task_failure(task_id, asyncio.CancelledError())

        interrupted = service.get_task(task_id)
        assert interrupted is not None
        assert interrupted.status == "running"
        assert interrupted.current_stage_name == "presentation_delivery"

        monkeypatch.setattr(agent_module, "generate_image", original_generate_image)
        fresh_service = ResearchAgentService()
        resumed = await fresh_service.continue_task(task_id)
        completed = fresh_service.get_task(task_id)

        assert resumed["status"] == "completed"
        assert completed is not None
        assert completed.status == "completed"
        assert Path(completed.artifact_root, "presentation", "STAGE_REPORT.pptx").exists()

    asyncio.run(scenario())
